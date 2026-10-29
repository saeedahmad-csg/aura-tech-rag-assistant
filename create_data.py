"""
Script to generate data/products.xlsx for Aura Technologies Inc.
"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

os.makedirs("data", exist_ok=True)
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Product Catalog & Pricing"

headers = [
    "Product Name", "SKU", "Tier", "Monthly Price ($)", 
    "Annual Price ($/yr)", "Storage Included", "Max Users", 
    "SLA Guarantee", "Support Level", "Key Features"
]

rows = [
    [
        "Aura Cloud Starter", "AC-101", "Starter", 49, 490, 
        "100 GB", "Up to 5 users", "99.5%", "Standard Email (24h response)",
        "Core cloud storage, basic REST API access, role-based permissions, automated daily backups."
    ],
    [
        "Aura Cloud Professional", "AC-201", "Professional", 149, 1490, 
        "1 TB", "Up to 25 users", "99.9%", "Priority Email & Chat (4h response)",
        "Everything in Starter plus advanced analytics dashboard, custom webhooks, audit logs, SSO with Google & Okta."
    ],
    [
        "Aura Cloud Enterprise", "AC-301", "Enterprise", 499, 4990, 
        "5 TB", "Unlimited users", "99.99%", "24/7 Phone & Slack Support (15m response)",
        "Everything in Professional plus dedicated database instance, custom AI agent integrations, HIPAA and SOC-2 compliance pack, custom domain."
    ],
    [
        "Aura AI Studio Add-on", "AI-501", "Add-on", 99, 990, 
        "500 GB Vector Store", "All licensed users", "99.9%", "Matches parent tier",
        "Semantic search indexing, automated RAG pipelines, custom LLM fine-tuning connector, vector similarity analytics."
    ],
    [
        "Aura Dedicated Private Cloud", "AC-901", "Custom / On-Prem", 1999, 19990, 
        "Unlimited NVMe", "Unlimited", "99.999%", "Dedicated Technical Account Manager",
        "Air-gapped deployment option, dedicated single-tenant Kubernetes cluster, custom SLA, sovereign data residency, priority hardware upgrades."
    ]
]

# Write headers
ws.append(headers)

# Write rows
for row in rows:
    ws.append(row)

# Styling
header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

for col_num in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for row in ws.iter_rows(min_row=2, max_row=len(rows) + 1, min_col=1, max_col=len(headers)):
    for cell in row:
        cell.border = thin_border
        cell.alignment = Alignment(vertical="center")

# Adjust column widths
for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = openpyxl.utils.get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

excel_path = os.path.join("data", "products.xlsx")
wb.save(excel_path)
print(f"Successfully generated {excel_path} with {len(rows)} products.")

